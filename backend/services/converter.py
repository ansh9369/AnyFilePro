import os
import zipfile
from typing import List, Tuple
from PIL import Image
from pypdf import PdfReader, PdfWriter
import pypdfium2 as pdfium
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.utils import ImageReader
try:
    from docx import Document
except ImportError:
    Document = None
try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None
try:
    from pptx import Presentation
except ImportError:
    Presentation = None
try:
    from pdf2docx import Converter as DocxConverter
except ImportError:
    DocxConverter = None

from backend.core.config import settings

class ConverterService:

    @staticmethod
    def images_to_pdf(image_paths: List[str], output_filename: str) -> str:
        if not image_paths:
            raise ValueError("No images provided")

        images = []
        for path in image_paths:
            img = Image.open(path)
            if img.mode != 'RGB':
                img = img.convert('RGB')
            images.append(img)

        output_path = os.path.join(settings.TEMP_DIR, output_filename)
        images[0].save(output_path, save_all=True, append_images=images[1:])
        return output_path

    @staticmethod
    def pdf_to_images(pdf_path: str, output_base_name: str) -> str:
        pdf = pdfium.PdfDocument(pdf_path)
        n_pages = len(pdf)

        # If single page, return just the image
        if n_pages == 1:
            page = pdf[0]
            bitmap = page.render(scale=2) # scale=2 for better quality
            pil_image = bitmap.to_pil()
            output_path = os.path.join(settings.TEMP_DIR, f"{output_base_name}.png")
            pil_image.save(output_path)
            return output_path

        # If multiple, return a zip
        zip_filename = f"{output_base_name}.zip"
        zip_path = os.path.join(settings.TEMP_DIR, zip_filename)

        with zipfile.ZipFile(zip_path, 'w') as zipf:
            for i in range(n_pages):
                page = pdf[i]
                bitmap = page.render(scale=2)
                pil_image = bitmap.to_pil()
                img_name = f"{output_base_name}_page_{i+1}.png"
                img_path = os.path.join(settings.TEMP_DIR, img_name)
                pil_image.save(img_path)
                zipf.write(img_path, arcname=img_name)
                os.remove(img_path) # Clean up individual files

        return zip_path

    @staticmethod
    def merge_pdfs(pdf_paths: List[str], output_filename: str) -> str:
        merger = PdfWriter()
        for path in pdf_paths:
            merger.append(path)

        output_path = os.path.join(settings.TEMP_DIR, output_filename)
        merger.write(output_path)
        merger.close()
        return output_path

    @staticmethod
    def split_pdf(pdf_path: str, output_base_name: str) -> str:
        reader = PdfReader(pdf_path)

        zip_filename = f"{output_base_name}_split.zip"
        zip_path = os.path.join(settings.TEMP_DIR, zip_filename)

        with zipfile.ZipFile(zip_path, 'w') as zipf:
            for i, page in enumerate(reader.pages):
                writer = PdfWriter()
                writer.add_page(page)

                split_name = f"{output_base_name}_page_{i+1}.pdf"
                split_path = os.path.join(settings.TEMP_DIR, split_name)

                with open(split_path, "wb") as f:
                    writer.write(f)

                zipf.write(split_path, arcname=split_name)
                os.remove(split_path)

        return zip_path

    @staticmethod
    def rotate_pdf(pdf_path: str, degrees: int, output_filename: str) -> str:
        reader = PdfReader(pdf_path)
        writer = PdfWriter()

        for page in reader.pages:
            writer.add_page(page)
            writer.pages[-1].rotate(degrees)

        output_path = os.path.join(settings.TEMP_DIR, output_filename)
        with open(output_path, "wb") as f:
            writer.write(f)
        return output_path

    @staticmethod
    def docx_to_pdf(docx_path: str, output_filename: str) -> str:
        if not Document:
            raise ImportError("python-docx not installed")

        doc = Document(docx_path)
        output_path = os.path.join(settings.TEMP_DIR, output_filename)
        c = canvas.Canvas(output_path, pagesize=letter)
        width, height = letter
        y = height - 40

        for para in doc.paragraphs:
            text = para.text
            if y < 40:
                c.showPage()
                y = height - 40

            # Simple text wrapping could be added here
            c.drawString(40, y, text[:100]) # Truncate for MVP stability
            y -= 14

        c.save()
        return output_path

    @staticmethod
    def xlsx_to_pdf(xlsx_path: str, output_filename: str) -> str:
        if not load_workbook:
            raise ImportError("openpyxl not installed")

        wb = load_workbook(xlsx_path)
        ws = wb.active
        output_path = os.path.join(settings.TEMP_DIR, output_filename)
        c = canvas.Canvas(output_path, pagesize=letter)
        width, height = letter
        y = height - 40
        x = 40

        for row in ws.iter_rows(values_only=True):
            line = " | ".join([str(cell) if cell else "" for cell in row])
            if y < 40:
                c.showPage()
                y = height - 40
            c.drawString(x, y, line[:100])
            y -= 14

        c.save()
        return output_path

    @staticmethod
    def pptx_to_pdf(pptx_path: str, output_filename: str) -> str:
        if not Presentation:
            raise ImportError("python-pptx not installed")

        prs = Presentation(pptx_path)
        output_path = os.path.join(settings.TEMP_DIR, output_filename)
        c = canvas.Canvas(output_path, pagesize=letter)
        width, height = letter

        for slide in prs.slides:
            y = height - 40
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text
                    if y < 40:
                        break # Clip content for MVP
                    c.drawString(40, y, text[:100])
                    y -= 20
            c.showPage()

        c.save()
        return output_path

    @staticmethod
    def pdf_to_docx(pdf_path: str, output_filename: str) -> str:
        if not DocxConverter:
             raise ImportError("pdf2docx not installed")

        output_path = os.path.join(settings.TEMP_DIR, output_filename)
        cv = DocxConverter(pdf_path)
        cv.convert(output_path)
        cv.close()
        return output_path

    @staticmethod
    def compress_pdf(pdf_path: str, output_filename: str) -> str:
        reader = PdfReader(pdf_path)
        writer = PdfWriter()

        for page in reader.pages:
            writer.add_page(page)

        for page in writer.pages:
            page.compress_content_streams()  # This provides lossless compression

        output_path = os.path.join(settings.TEMP_DIR, output_filename)
        with open(output_path, "wb") as f:
            writer.write(f)
        return output_path

    @staticmethod
    def lock_pdf(pdf_path: str, password: str, output_filename: str) -> str:
        reader = PdfReader(pdf_path)
        writer = PdfWriter()
        writer.append_pages_from_reader(reader)
        writer.encrypt(password)

        output_path = os.path.join(settings.TEMP_DIR, output_filename)
        with open(output_path, "wb") as f:
            writer.write(f)
        return output_path

    @staticmethod
    def unlock_pdf(pdf_path: str, password: str, output_filename: str) -> str:
        reader = PdfReader(pdf_path)
        if reader.is_encrypted:
            reader.decrypt(password)

        writer = PdfWriter()
        writer.append_pages_from_reader(reader)

        output_path = os.path.join(settings.TEMP_DIR, output_filename)
        with open(output_path, "wb") as f:
            writer.write(f)
        return output_path
